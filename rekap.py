"""Rekap daftar penugasan ke Excel (openpyxl), atau CSV kalau openpyxl tidak ada."""
from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from core import OUTPUT, ST, Penugasan, format_tanggal

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

KOLOM = [
    # (judul, lebar, fungsi nilai)
    ("No", 5, None),
    ("Nomor ST", 24, lambda p: p.nomor_st),
    ("Tanggal ST", 12, lambda p: p.tanggal_st),
    ("Jenis", 12, lambda p: p.jenis),
    ("Judul", 46, lambda p: p.judul),
    ("Objek", 24, lambda p: p.objek),
    ("Peran Saya", 18, lambda p: p.peran_saya),
    ("Tim", 40, lambda p: "; ".join(f"{a['nama']} ({a['peran']})" for a in p.tim if a.get("nama"))),
    ("Mulai", 12, lambda p: p.mulai),
    ("Sampai", 12, lambda p: p.sampai),
    ("Jumlah Hari", 8, lambda p: p.jumlah_hari),
    ("Tenggat", 12, lambda p: p.tenggat),
    ("Tahap", 15, lambda p: p.tahap),
    ("Status", 13, lambda p: p.status_tampil()),
    ("Tanggal Selesai", 12, lambda p: p.selesai_pada),
    ("Surat Tugas", 8, lambda p: "Ada" if p.punya(ST) else "-"),
    ("Output", 40, lambda p: "; ".join(b.name for b in p.berkas(OUTPUT)) or "-"),
    ("Catatan", 30, lambda p: p.catatan),
    ("Folder", 16, None),
]


def simpan(daftar: list[Penugasan], path: Path, judul: str = "Rekap Penugasan") -> Path:
    """Tulis rekap ke `path`. Kalau openpyxl tidak ada, akhiran diganti .csv."""
    path = Path(path)
    if Workbook is None:
        path = path.with_suffix(".csv")
        _csv(daftar, path)
    else:
        _xlsx(daftar, path, judul)
    return path


def _nilai(p: Penugasan, i: int, fn):
    if fn is None:
        return i
    return fn(p)


def _csv(daftar: list[Penugasan], path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # BOM: Excel membaca UTF-8 dengan benar
        w = csv.writer(f, delimiter=";")
        w.writerow([k for k, _, _ in KOLOM])
        for i, p in enumerate(daftar, 1):
            baris = []
            for judul, _, fn in KOLOM:
                v = str(p.folder) if judul == "Folder" else _nilai(p, i, fn)
                baris.append(format_tanggal(v) if isinstance(v, date) else ("" if v is None else v))
            w.writerow(baris)


def _xlsx(daftar: list[Penugasan], path: Path, judul: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap"
    tipis = Side(style="thin", color="BFC8D6")
    garis = Border(left=tipis, right=tipis, top=tipis, bottom=tipis)

    ws["A1"] = judul
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Dibuat {format_tanggal(date.today())} · {len(daftar)} penugasan"
    ws["A2"].font = Font(italic=True, color="666666")

    baris_kepala = 4
    for c, (nama, lebar, _) in enumerate(KOLOM, 1):
        sel = ws.cell(row=baris_kepala, column=c, value=nama)
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = PatternFill("solid", fgColor="1A6F96")
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sel.border = garis
        ws.column_dimensions[get_column_letter(c)].width = lebar

    for i, p in enumerate(daftar, 1):
        r = baris_kepala + i
        for c, (nama, _, fn) in enumerate(KOLOM, 1):
            if nama == "Folder":
                sel = ws.cell(row=r, column=c, value="Buka folder")
                if p.folder:
                    sel.hyperlink = str(p.folder)
                    sel.font = Font(color="1A6F96", underline="single")
            else:
                sel = ws.cell(row=r, column=c, value=_nilai(p, i, fn))
                if isinstance(sel.value, date):
                    sel.number_format = "DD-MM-YYYY"
            sel.border = garis
            sel.alignment = Alignment(vertical="top", wrap_text=nama in ("Judul", "Tim", "Output", "Catatan"))

    ws.freeze_panes = ws.cell(row=baris_kepala + 1, column=3)
    if daftar:
        ws.auto_filter.ref = f"A{baris_kepala}:{get_column_letter(len(KOLOM))}{baris_kepala + len(daftar)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    wb.save(path)
